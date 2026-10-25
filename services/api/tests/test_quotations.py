import pytest
from httpx import AsyncClient
from decimal import Decimal
from datetime import date, timedelta


@pytest.mark.asyncio
async def test_quotation_lifecycle_and_immutable_snapshot(client: AsyncClient, token_a: str, token_b: str):
    headers_a = {"Authorization": f"Bearer {token_a}"}
    headers_b = {"Authorization": f"Bearer {token_b}"}

    # 1. Setup Org A and Org B
    res_a = await client.post("/api/organizations", json={"name": "Larsen & Toubro Ltd"}, headers=headers_a)
    assert res_a.status_code in (200, 201)
    org_a_id = res_a.json()["id"]

    res_b = await client.post("/api/organizations", json={"name": "Shapoorji Pallonji Co"}, headers=headers_b)
    assert res_b.status_code in (200, 201)
    org_b_id = res_b.json()["id"]

    # 2. Create Client and Project in Org A
    c_res = await client.post(
        f"/api/organizations/{org_a_id}/clients",
        json={"name": "Adani Realty", "billing_address": "BKC, Mumbai", "gstin": "27AAACA0000A1Z5"},
        headers=headers_a
    )
    assert c_res.status_code == 201
    client_a_id = c_res.json()["id"]

    p_res = await client.post(
        f"/api/organizations/{org_a_id}/projects",
        json={"name": "Marine Drive Tower", "project_type": "building", "client_id": client_a_id},
        headers=headers_a
    )
    assert p_res.status_code == 201
    proj_a_id = p_res.json()["id"]

    # 3. Create Estimate and BOQ Items in Org A
    est_res = await client.post(
        f"/api/organizations/{org_a_id}/estimates",
        json={
            "name": "Tower Foundation & Civil",
            "project_id": proj_a_id,
            "project_type": "building",
            "overhead_percent": 10.0,
            "contingency_percent": 5.0,
            "discount_percent": 2.0,
            "gst_percent": 0.0,
        },
        headers=headers_a
    )
    assert est_res.status_code == 201
    est_a_id = est_res.json()["id"]

    # Add 2 BOQ items
    await client.post(
        f"/api/organizations/{org_a_id}/estimates/{est_a_id}/items",
        json={"item_number": "1.1", "description": "Pile Foundation Excavation", "quantity": 100.0, "unit_rate": 500.0, "unit": "cum"},
        headers=headers_a
    )
    await client.post(
        f"/api/organizations/{org_a_id}/estimates/{est_a_id}/items",
        json={"item_number": "1.2", "description": "M30 Grade Reinforced Concrete", "quantity": 50.0, "unit_rate": 6000.0, "unit": "cum"},
        headers=headers_a
    )
    # Estimate subtotal = (100 * 500) + (50 * 6000) = 50,000 + 300,000 = 350,000 INR

    # 4. Create Draft Quotation from Estimate
    q_create_res = await client.post(
        f"/api/organizations/{org_a_id}/quotations",
        json={
            "estimate_id": est_a_id,
            "language": "mr",
            "valid_until_date": (date.today() + timedelta(days=30)).isoformat(),
            "payment_terms": "२०% आगाऊ, ६०% कामाच्या प्रगतीनुसार, २०% पूर्ण झाल्यावर.",
            "exclusions": "वीज आणि पाण्याची सोय ग्राहकाने करावी.",
            "notes": "दरपत्रक ३० दिवसांसाठी वैध.",
            "tax_configured": False,
        },
        headers=headers_a
    )
    assert q_create_res.status_code == 201
    q_data = q_create_res.json()
    q_id = q_data["id"]
    assert q_data["status"] == "draft"
    assert q_data["language"] == "mr"
    assert q_data["quotation_number"].startswith("QT-")
    assert float(q_data["subtotal"]) == 350000.0
    assert q_data["tax_configured"] is False
    assert q_data["snapshot"] is None

    # 5. Edit Draft Quotation
    q_patch_res = await client.patch(
        f"/api/organizations/{org_a_id}/quotations/{q_id}",
        json={
            "notes": "सुधारित टीप: दरपत्रक ४५ दिवसांसाठी वैध आहे.",
            "authorized_signatory_name": "विक्रम साराभाई",
            "authorized_signatory_designation": "मुख्य प्रकल्प व्यवस्थापक"
        },
        headers=headers_a
    )
    assert q_patch_res.status_code == 200
    assert "४५ दिवस" in q_patch_res.json()["notes"]
    assert q_patch_res.json()["authorized_signatory_name"] == "विक्रम साराभाई"

    # 6. Issue Quotation: Verify unconfigured tax rejection unless confirmed
    # Attempting to issue without tax configuration/confirmation should fail
    fail_issue = await client.post(
        f"/api/organizations/{org_a_id}/quotations/{q_id}/issue",
        json={"confirm_tax_configured": False},
        headers=headers_a
    )
    assert fail_issue.status_code == 400
    assert "tax treatment has not been configured" in fail_issue.json()["detail"]

    # Issue with explicit 18% GST confirmation
    issue_res = await client.post(
        f"/api/organizations/{org_a_id}/quotations/{q_id}/issue",
        json={
            "confirm_tax_configured": True,
            "tax_name": "GST",
            "tax_rate": 18.0
        },
        headers=headers_a
    )
    assert issue_res.status_code == 200
    issued_data = issue_res.json()
    assert issued_data["status"] == "issued"
    assert issued_data["tax_configured"] is True
    assert issued_data["snapshot"] is not None
    snap = issued_data["snapshot"]
    assert snap["quotation_number"] == issued_data["quotation_number"]
    assert len(snap["items"]) == 2
    assert float(snap["financials"]["subtotal"]) == 350000.0
    assert float(snap["financials"]["tax_rate"]) == 18.0

    # 7. CRITICAL TEST: Snapshot Immutability
    # Modify the original estimate and BOQ items after quotation was issued!
    await client.post(
        f"/api/organizations/{org_a_id}/estimates/{est_a_id}/items",
        json={"item_number": "1.3", "description": "Extra Unrelated Item", "quantity": 1000.0, "unit_rate": 5000.0},
        headers=headers_a
    )
    # Also change client name
    await client.patch(
        f"/api/organizations/{org_a_id}/clients/{client_a_id}",
        json={"name": "Changed Client Name"},
        headers=headers_a
    )

    # Fetch the issued quotation again: it MUST retain the immutable snapshot!
    get_issued = await client.get(
        f"/api/organizations/{org_a_id}/quotations/{q_id}",
        headers=headers_a
    )
    assert get_issued.status_code == 200
    issued_check = get_issued.json()
    # Number of items on quotation MUST still be 2 (not 3!)
    assert len(issued_check["items"]) == 2
    assert float(issued_check["subtotal"]) == 350000.0
    assert issued_check["client_details"]["name"] == "Adani Realty"  # Preserved original client!

    # 8. Editing an issued quotation must be strictly rejected
    edit_issued_fail = await client.patch(
        f"/api/organizations/{org_a_id}/quotations/{q_id}",
        json={"notes": "Illegal edit attempt"},
        headers=headers_a
    )
    assert edit_issued_fail.status_code == 400
    assert "Cannot edit quotation in 'issued' status" in edit_issued_fail.json()["detail"]

    # 9. Status Transitions
    # 'issued' -> 'accepted' (valid)
    accept_res = await client.patch(
        f"/api/organizations/{org_a_id}/quotations/{q_id}/status",
        json={"status": "accepted"},
        headers=headers_a
    )
    assert accept_res.status_code == 200
    assert accept_res.json()["status"] == "accepted"

    # 'accepted' -> 'draft' (invalid!)
    invalid_trans = await client.patch(
        f"/api/organizations/{org_a_id}/quotations/{q_id}/status",
        json={"status": "draft"},
        headers=headers_a
    )
    assert invalid_trans.status_code in (400, 422)

    # 10. Duplicate Action
    dup_res = await client.post(
        f"/api/organizations/{org_a_id}/quotations/{q_id}/duplicate",
        headers=headers_a
    )
    assert dup_res.status_code == 201
    dup_data = dup_res.json()
    assert dup_data["status"] == "draft"
    assert dup_data["id"] != q_id
    assert dup_data["quotation_number"] != issued_data["quotation_number"]
    assert dup_data["snapshot"] is None

    # 11. PDF and Excel Exports
    pdf_res = await client.get(
        f"/api/organizations/{org_a_id}/quotations/{q_id}/pdf",
        headers=headers_a
    )
    assert pdf_res.status_code == 200
    assert pdf_res.headers["content-type"] == "application/pdf"
    assert pdf_res.content.startswith(b"%PDF")
    assert len(pdf_res.content) > 1000

    excel_res = await client.get(
        f"/api/organizations/{org_a_id}/quotations/{q_id}/excel",
        headers=headers_a
    )
    assert excel_res.status_code == 200
    assert "spreadsheetml" in excel_res.headers["content-type"]
    assert excel_res.content.startswith(b"PK\x03\x04")

    # 12. Cross-Organization Tenant Isolation
    # User B in Org B must be strictly denied access to Org A's quotation
    cross_get = await client.get(f"/api/organizations/{org_a_id}/quotations/{q_id}", headers=headers_b)
    assert cross_get.status_code in (403, 404)

    cross_issue = await client.post(f"/api/organizations/{org_a_id}/quotations/{q_id}/issue", json={}, headers=headers_b)
    assert cross_issue.status_code in (403, 404)

    cross_pdf = await client.get(f"/api/organizations/{org_a_id}/quotations/{q_id}/pdf", headers=headers_b)
    assert cross_pdf.status_code in (403, 404)

    cross_excel = await client.get(f"/api/organizations/{org_a_id}/quotations/{q_id}/excel", headers=headers_b)
    assert cross_excel.status_code in (403, 404)

    cross_dup = await client.post(f"/api/organizations/{org_a_id}/quotations/{q_id}/duplicate", headers=headers_b)
    assert cross_dup.status_code in (403, 404)


@pytest.mark.asyncio
async def test_quotation_search_filter_and_pagination(client: AsyncClient, token_a: str):
    headers_a = {"Authorization": f"Bearer {token_a}"}

    # Setup Org A
    res_a = await client.post("/api/organizations", json={"name": "Tata Projects Infrastructure"}, headers=headers_a)
    org_a_id = res_a.json()["id"]

    # Setup Client and Project
    c_res = await client.post(
        f"/api/organizations/{org_a_id}/clients",
        json={"name": "Lodha Developers"},
        headers=headers_a
    )
    client_id = c_res.json()["id"]

    p_res = await client.post(
        f"/api/organizations/{org_a_id}/projects",
        json={"name": "Lodha Park Mumbai", "project_type": "building", "client_id": client_id},
        headers=headers_a
    )
    proj_id = p_res.json()["id"]

    # Create 3 estimates and quotations
    for i in range(1, 4):
        est = await client.post(
            f"/api/organizations/{org_a_id}/estimates",
            json={"name": f"Wing {i} Estimate", "project_id": proj_id, "project_type": "building"},
            headers=headers_a
        )
        est_id = est.json()["id"]
        await client.post(
            f"/api/organizations/{org_a_id}/estimates/{est_id}/items",
            json={"item_number": "1.1", "description": f"Civil work wing {i}", "quantity": 10.0 * i, "unit_rate": 1000.0},
            headers=headers_a
        )
        q = await client.post(
            f"/api/organizations/{org_a_id}/quotations",
            json={"estimate_id": est_id, "language": "hi" if i == 1 else "en"},
            headers=headers_a
        )
        assert q.status_code == 201

    # Test listing
    list_all = await client.get(f"/api/organizations/{org_a_id}/quotations", headers=headers_a)
    assert list_all.status_code == 200
    assert len(list_all.json()) == 3

    # Test pagination
    list_paginated = await client.get(f"/api/organizations/{org_a_id}/quotations?skip=0&limit=2", headers=headers_a)
    assert len(list_paginated.json()) == 2

    # Test search by client
    list_search = await client.get(f"/api/organizations/{org_a_id}/quotations?search=Lodha", headers=headers_a)
    assert len(list_search.json()) == 3

    # Test filter by status
    list_draft = await client.get(f"/api/organizations/{org_a_id}/quotations?status=draft", headers=headers_a)
    assert len(list_draft.json()) == 3

    list_issued = await client.get(f"/api/organizations/{org_a_id}/quotations?status=issued", headers=headers_a)
    assert len(list_issued.json()) == 0


@pytest.mark.asyncio
async def test_quotation_security_and_edge_cases(client: AsyncClient, token_a: str, token_b: str):
    headers_a = {"Authorization": f"Bearer {token_a}"}
    headers_b = {"Authorization": f"Bearer {token_b}"}

    # Setup Org A and Org B
    res_a = await client.post("/api/organizations", json={"name": "Tata Projects Ltd"}, headers=headers_a)
    org_a_id = res_a.json()["id"]

    res_b = await client.post("/api/organizations", json={"name": "NCC Urban Ltd"}, headers=headers_b)
    org_b_id = res_b.json()["id"]

    # Create project & estimate in Org A
    p_res = await client.post(f"/api/organizations/{org_a_id}/projects", json={"name": "Metro Viaduct"}, headers=headers_a)
    proj_a_id = p_res.json()["id"]

    est_res = await client.post(
        f"/api/organizations/{org_a_id}/estimates",
        json={"name": "Viaduct Pier 1-10", "project_id": proj_a_id, "project_type": "road"},
        headers=headers_a
    )
    est_a_id = est_res.json()["id"]

    # 1. Non-existent estimate ID returns 404
    bad_est = await client.post(
        f"/api/organizations/{org_a_id}/quotations",
        json={"estimate_id": "00000000-0000-0000-0000-000000000000"},
        headers=headers_a
    )
    assert bad_est.status_code == 404

    # 2. Cross-org estimate theft: Org B tries to create quotation using Org A estimate
    theft_res = await client.post(
        f"/api/organizations/{org_b_id}/quotations",
        json={"estimate_id": est_a_id},
        headers=headers_b
    )
    assert theft_res.status_code == 404

    # 3. Create valid quotation in Org A
    q_res = await client.post(
        f"/api/organizations/{org_a_id}/quotations",
        json={"estimate_id": est_a_id, "quotation_number": "QT-SEC-001"},
        headers=headers_a
    )
    assert q_res.status_code == 201
    quote_id = q_res.json()["id"]

    # 4. Duplicate quotation number in same org rejected with 400
    dup_num_res = await client.post(
        f"/api/organizations/{org_a_id}/quotations",
        json={"estimate_id": est_a_id, "quotation_number": "QT-SEC-001"},
        headers=headers_a
    )
    assert dup_num_res.status_code == 400
    assert "already exists" in dup_num_res.json()["detail"]

    # 5. Invalid tax rate (negative or >100) rejected with 422
    neg_tax = await client.post(
        f"/api/organizations/{org_a_id}/quotations",
        json={"estimate_id": est_a_id, "tax_rate": -18.0},
        headers=headers_a
    )
    assert neg_tax.status_code == 422

    over_tax = await client.post(
        f"/api/organizations/{org_a_id}/quotations",
        json={"estimate_id": est_a_id, "tax_rate": 150.0},
        headers=headers_a
    )
    assert over_tax.status_code == 422

    # 6. Status transition: Cannot transition draft directly to accepted (must issue first)
    bad_trans = await client.patch(
        f"/api/organizations/{org_a_id}/quotations/{quote_id}/status",
        json={"status": "accepted"},
        headers=headers_a
    )
    assert bad_trans.status_code == 400

    # 7. Non-existent quotation ID returns 404
    missing_q = await client.get(
        f"/api/organizations/{org_a_id}/quotations/00000000-0000-0000-0000-000000000000",
        headers=headers_a
    )
    assert missing_q.status_code == 404

    # 8. Cross-org denial: Org B cannot read Org A quotation
    cross_read = await client.get(
        f"/api/organizations/{org_a_id}/quotations/{quote_id}",
        headers=headers_b
    )
    assert cross_read.status_code == 403

    # Org B cannot access Org A quotation even if passing Org B org_id
    cross_spoof = await client.get(
        f"/api/organizations/{org_b_id}/quotations/{quote_id}",
        headers=headers_b
    )
    assert cross_spoof.status_code == 404

    # 9. Issue quotation successfully
    issue_res = await client.post(
        f"/api/organizations/{org_a_id}/quotations/{quote_id}/issue",
        json={"confirm_tax_configured": True, "tax_rate": 18.0, "tax_name": "GST 18%"},
        headers=headers_a
    )
    assert issue_res.status_code == 200
    assert issue_res.json()["status"] == "issued"

    # 10. Re-issuing an already issued quotation rejected with 400
    reissue_res = await client.post(
        f"/api/organizations/{org_a_id}/quotations/{quote_id}/issue",
        json={"confirm_tax_configured": True},
        headers=headers_a
    )
    assert reissue_res.status_code == 400

    # 11. Editing an issued quotation rejected with 400
    edit_issued = await client.patch(
        f"/api/organizations/{org_a_id}/quotations/{quote_id}",
        json={"payment_terms": "Modified after issue"},
        headers=headers_a
    )
    assert edit_issued.status_code == 400

    # 12. Valid status transition: issued -> accepted
    accept_res = await client.patch(
        f"/api/organizations/{org_a_id}/quotations/{quote_id}/status",
        json={"status": "accepted"},
        headers=headers_a
    )
    assert accept_res.status_code == 200
    assert accept_res.json()["status"] == "accepted"


@pytest.mark.asyncio
async def test_excel_formula_injection_and_multilingual_pdf(client: AsyncClient, token_a: str):
    import io
    from openpyxl import load_workbook

    headers_a = {"Authorization": f"Bearer {token_a}"}

    # Setup Org
    res_a = await client.post("/api/organizations", json={"name": "Godrej Construction"}, headers=headers_a)
    org_id = res_a.json()["id"]

    p_res = await client.post(f"/api/organizations/{org_id}/projects", json={"name": "Vikhroli IT Park"}, headers=headers_a)
    proj_id = p_res.json()["id"]

    est_res = await client.post(
        f"/api/organizations/{org_id}/estimates",
        json={"name": "Civil Package", "project_id": proj_id, "project_type": "building"},
        headers=headers_a
    )
    est_id = est_res.json()["id"]

    # Add an item with potential formula injection strings
    await client.post(
        f"/api/organizations/{org_id}/estimates/{est_id}/items",
        json={
            "item_number": "1.1",
            "description": "=1+1 Formula Injection Attempt",
            "quantity": 10.0,
            "unit_rate": 500.0,
            "unit": "cum",
            "notes": "@SUM(A1:B1) malicious note"
        },
        headers=headers_a
    )

    # Create quotation in Marathi
    q_res = await client.post(
        f"/api/organizations/{org_id}/quotations",
        json={
            "estimate_id": est_id,
            "language": "mr",
            "notes": "-cmd.exe malicious command trigger",
            "payment_terms": "+30% advance on signing"
        },
        headers=headers_a
    )
    assert q_res.status_code == 201
    quote_id = q_res.json()["id"]

    # Issue quotation
    await client.post(
        f"/api/organizations/{org_id}/quotations/{quote_id}/issue",
        json={"confirm_tax_configured": True, "tax_rate": 18.0},
        headers=headers_a
    )

    # 1. Test Excel Export and Formula Sanitization
    excel_res = await client.get(f"/api/organizations/{org_id}/quotations/{quote_id}/excel", headers=headers_a)
    assert excel_res.status_code == 200
    assert len(excel_res.content) > 1000

    # Parse with openpyxl and inspect all string cells
    wb = load_workbook(io.BytesIO(excel_res.content))
    ws = wb.active
    for row in ws.iter_rows(values_only=True):
        for cell_val in row:
            if isinstance(cell_val, str):
                # Verify that no cell starts directly with active formula triggers without apostrophe
                assert not cell_val.startswith("="), f"Unsanitized formula cell found: {cell_val}"
                assert not (cell_val.startswith(("+", "-", "@")) and len(cell_val) > 1 and cell_val[1].isalpha()), f"Unsanitized cell: {cell_val}"

    # 2. Test Multilingual PDF Generation (Marathi with Devanagari shaping)
    pdf_res = await client.get(f"/api/organizations/{org_id}/quotations/{quote_id}/pdf", headers=headers_a)
    assert pdf_res.status_code == 200
    assert pdf_res.headers["content-type"] == "application/pdf"
    assert pdf_res.content.startswith(b"%PDF-")
    assert len(pdf_res.content) > 2000
